"""Generate tests/fixtures/import/sample.docx (python-docx). Run: python3 scripts/fixtures/make_docx.py"""
import io, struct, zlib, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT

def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

d = Document()
d.add_heading('HandNote import test', 0)
d.add_heading('1. Headings and paragraphs', level=1)
p = d.add_paragraph('A normal paragraph with ')
r = p.add_run('bold'); r.bold = True
p.add_run(', ')
r = p.add_run('italic'); r.italic = True
p.add_run(', ')
r = p.add_run('underlined'); r.underline = True
p.add_run(' and ')
r = p.add_run('red'); r.font.color.rgb = RGBColor(0xC0, 0x10, 0x10)
p.add_run(' text. ' + 'This sentence is repeated to wrap over several lines. ' * 4)
p = d.add_paragraph('Centered paragraph.'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = d.add_paragraph('Right aligned paragraph.'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
d.add_paragraph('中文段落：手写笔记插件把文档导入为固定页面，批注不会因为字体不同而错位。' * 3)
d.add_heading('2. Lists', level=1)
for s in ['First bullet', 'Second bullet', 'Third bullet']:
    d.add_paragraph(s, style='List Bullet')
for s in ['Step one', 'Step two', 'Step three']:
    d.add_paragraph(s, style='List Number')
d.add_heading('3. Table', level=1)
t = d.add_table(rows=3, cols=3)
t.style = 'Table Grid'
for i in range(3):
    for j in range(3):
        t.cell(i, j).text = f'R{i+1}C{j+1}'
a = t.cell(2, 0).merge(t.cell(2, 1)); a.text = 'merged R3C1-2'
d.add_heading('4. Picture', level=1)
d.add_picture(io.BytesIO(png(120, 60, (30, 120, 220))), width=Inches(2))
p = d.add_paragraph('After the picture.')
p.add_run().add_break(WD_BREAK.PAGE)
d.add_heading('5. Second page', level=1)
d.add_paragraph('This paragraph starts on page 2 after a manual page break.')
sec = d.add_section()
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
d.add_heading('6. Landscape section', level=1)
d.add_paragraph('This page is landscape.')
out = sys.argv[1] if len(sys.argv) > 1 else 'tests/fixtures/import/sample.docx'
d.save(out)
print(out)
