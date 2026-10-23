"""Generate tests/fixtures/import/sample.pptx (python-pptx). Run: python3 scripts/fixtures/make_pptx.py"""
import io, struct, zlib, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
# 1 title slide
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'HandNote slides test'
s.placeholders[1].text = 'Subtitle: one page per slide'
# 2 bullets
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Bullets and levels'
tf = s.placeholders[1].text_frame
tf.text = 'First point'
for t, lvl in [('Second point', 0), ('Sub point A', 1), ('Sub point B', 1), ('中文要点：幻灯片导入', 0)]:
    p = tf.add_paragraph(); p.text = t; p.level = lvl
# 3 picture, shapes, background colour
s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(0xFF, 0xF4, 0xD6)
s.shapes.add_picture(io.BytesIO(png(300, 150, (30, 120, 220))), Inches(1), Inches(1), Inches(4), Inches(2))
r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(1), Inches(3), Inches(1.5))
r.fill.solid(); r.fill.fore_color.rgb = RGBColor(0xC0, 0x10, 0x10); r.line.color.rgb = RGBColor(0, 0, 0)
r.text_frame.text = 'Red box'
e = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10), Inches(1), Inches(2), Inches(2))
e.fill.solid(); e.fill.fore_color.rgb = RGBColor(0x10, 0xA0, 0x40)
rr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(4), Inches(3), Inches(1.5))
rr.text_frame.text = 'Theme-styled shape'
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5), Inches(4.5), Inches(9), Inches(6))
c.line.color.rgb = RGBColor(0x20, 0x20, 0x20); c.line.width = Pt(3)
tb = s.shapes.add_textbox(Inches(9.5), Inches(4), Inches(3), Inches(1))
tb.text_frame.text = 'Text box, right'
tb.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
tb.text_frame.paragraphs[0].runs[0].font.bold = True
# 4 table
s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
s.shapes.title.text = 'Table'
t = s.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(8), Inches(2)).table
for i in range(3):
    for j in range(3):
        t.cell(i, j).text = f'R{i+1}C{j+1}'
out = sys.argv[1] if len(sys.argv) > 1 else 'tests/fixtures/import/sample.pptx'
prs.save(out)
print(out)

# ---- tests/fixtures/import/rich.pptx: Chinese bullets, a group of preset shapes, a cropped and a
# rotated picture, a bottom-anchored wrapping text box. rich-libreoffice.pptx is the same deck after
# a LibreOffice round trip (pptx → odp → pptx), i.e. written by another producer.
def png2(w, h, left, right):
    rows = []
    for _ in range(h):
        rows.append(b'\x00' + b''.join(bytes(right if x > w // 2 else left) for x in range(w)))
    raw = b''.join(rows)
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

from pptx.enum.text import MSO_ANCHOR
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '第二讲 线性方程组'
tf = s.placeholders[1].text_frame
tf.text = '高斯消元法的基本步骤'
for t, l in [('化为行阶梯形', 1), ('回代求解', 1), ('解的存在性与唯一性', 0)]:
    p = tf.add_paragraph(); p.text = t; p.level = l
s = prs.slides.add_slide(prs.slide_layouts[6])
g = s.shapes.add_group_shape()
for i, (c, shape) in enumerate([((200, 60, 60), MSO_SHAPE.ISOSCELES_TRIANGLE), ((60, 160, 60), MSO_SHAPE.DIAMOND), ((60, 60, 200), MSO_SHAPE.RIGHT_ARROW), ((200, 150, 0), MSO_SHAPE.STAR_5_POINT), ((120, 0, 160), MSO_SHAPE.HEXAGON), ((0, 140, 140), MSO_SHAPE.CHEVRON)]):
    sh = g.shapes.add_shape(shape, Inches(0.5 + i * 1.5), Inches(1), Inches(1.2), Inches(1.2))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*c)
pic = s.shapes.add_picture(io.BytesIO(png2(200, 100, (250, 0, 0), (0, 0, 250))), Inches(1), Inches(3.5), Inches(3), Inches(1.5))
pic.crop_right = 0.5
pic2 = s.shapes.add_picture(io.BytesIO(png(200, 100, (0, 200, 0))), Inches(5), Inches(3.5), Inches(2), Inches(1))
pic2.rotation = 30
tb = s.shapes.add_textbox(Inches(5), Inches(5.5), Inches(4), Inches(1.5))
tf = tb.text_frame; tf.word_wrap = True
tf.text = 'A long text box paragraph that wraps across several lines inside a narrow box on the slide.'
tf.vertical_anchor = MSO_ANCHOR.BOTTOM
if len(sys.argv) <= 1:
    prs.save('tests/fixtures/import/rich.pptx')
    print('tests/fixtures/import/rich.pptx (then: soffice --convert-to odp, --convert-to pptx → rich-libreoffice.pptx)')
