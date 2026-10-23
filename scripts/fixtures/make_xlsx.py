"""Generate tests/fixtures/import/sample.xlsx (openpyxl, then a LibreOffice round trip so formula
results are stored like Excel stores them). Run: python3 scripts/fixtures/make_xlsx.py"""
import io, os, struct, subprocess, sys, tempfile, zlib, shutil, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image

def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

wb = Workbook()
ws = wb.active
ws.title = '成绩表'
thin = Side(style='thin', color='000000')
ws['A1'] = '期中成绩统计'
ws['A1'].font = Font(bold=True, size=16, color='1F4E79')
ws.merge_cells('A1:E1')
ws['A1'].alignment = Alignment(horizontal='center')
head = ['姓名', '数学', '英语', '总分', '占比']
for i, h in enumerate(head):
    c = ws.cell(row=2, column=i + 1, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='4472C4')
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    c.alignment = Alignment(horizontal='center')
rows = [('张三', 92, 85), ('李四', 78.5, 91), ('王五', 88, 76)]
for r, (n, a, b) in enumerate(rows, start=3):
    ws.cell(row=r, column=1, value=n)
    ws.cell(row=r, column=2, value=a)
    ws.cell(row=r, column=3, value=b)
    ws.cell(row=r, column=4, value=f'=B{r}+C{r}')
    ws.cell(row=r, column=5, value=f'=D{r}/SUM($D$3:$D$5)').number_format = '0.0%'
    for k in range(1, 6):
        ws.cell(row=r, column=k).border = Border(left=thin, right=thin, top=thin, bottom=thin)
ws['A7'] = '日期'
ws['B7'] = datetime.date(2024, 10, 23)
ws['B7'].number_format = 'yyyy-mm-dd'
ws['A8'] = '金额'
ws['B8'] = 1234567.891
ws['B8'].number_format = '"¥"#,##0.00'
ws['A9'] = '太窄'
ws['B9'] = 123456789012345
ws['B9'].number_format = '#,##0'
ws['A10'] = 'This long text runs over the empty cells to its right'
ws['A11'] = 'Wrapped text inside one cell that is narrow'
ws['A11'].alignment = Alignment(wrap_text=True, vertical='top')
ws['B11'] = 'x'
ws.column_dimensions['A'].width = 16
ws.column_dimensions['E'].width = 10
ws.column_dimensions['B'].width = 17
ws.add_image(Image(io.BytesIO(png(120, 60, (30, 120, 220)))), 'G3')
# a second, long sheet (several pages)
ws2 = wb.create_sheet('Data')
for r in range(1, 151):
    for c in range(1, 15):
        ws2.cell(row=r, column=c, value=r * 100 + c)
hidden = wb.create_sheet('Secret')
hidden['A1'] = 'hidden sheet'
hidden.sheet_state = 'hidden'
out = sys.argv[1] if len(sys.argv) > 1 else 'tests/fixtures/import/sample.xlsx'
tmp = tempfile.mkdtemp()
src = os.path.join(tmp, 'raw.xlsx')
wb.save(src)
# LibreOffice recalculates and stores the formula results (openpyxl stores none)
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', os.path.join(tmp, 'out'), src], check=True, capture_output=True)
shutil.copy(os.path.join(tmp, 'out', 'raw.xlsx'), out)
print(out)
